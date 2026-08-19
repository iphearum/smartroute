"""Parse and validate bulk product spreadsheets (CSV / XLSX) for import into
a business catalog.

Pure parsing and validation -- no database or HTTP access -- so every column
mapping, coercion, and rejection rule below is unit testable against an
in-memory bytes payload. The caller (api/commerce.py) is responsible for
actually writing accepted rows.

Contract: one row per product variant. A product that already exists by name
within the business gains another variant; otherwise the product is created.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

MAX_ROWS = 2000
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

# Accepted spellings per logical column. Compared after lowercasing and
# collapsing separators, so "Compare At Price", "compare_at_price", and
# "compare-at-price" all resolve to the same field.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("name", "product", "productname", "item", "itemname", "title"),
    "sku": ("sku", "code", "productcode", "barcode"),
    "price": ("price", "unitprice", "sellprice", "amount"),
    "compare_at_price": ("compareatprice", "originalprice", "wasprice", "listprice"),
    "category": ("category", "type", "group"),
    "description": ("description", "details", "notes"),
    "quantity": ("quantity", "qty", "stock", "stockqty", "quantityavailable"),
    "variant_title": ("varianttitle", "variant", "option", "size"),
}
REQUIRED_COLUMNS = ("name", "price")


class ProductImportError(RuntimeError):
    """Raised when a file cannot be read at all (bad format, no header, empty)."""


@dataclass
class ImportRow:
    row_number: int
    name: str
    price: Decimal
    sku: str | None = None
    compare_at_price: Decimal | None = None
    category: str | None = None
    description: str | None = None
    quantity: int | None = None
    variant_title: str | None = None


@dataclass
class ImportPreview:
    columns: dict[str, str] = field(default_factory=dict)
    rows: list[ImportRow] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    ignored_columns: list[str] = field(default_factory=list)

    @property
    def valid_count(self) -> int:
        return len(self.rows)


def _normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").strip().lower())


def map_columns(headers: list[str]) -> tuple[dict[str, int], list[str]]:
    """Resolve spreadsheet headers to logical field names.

    Returns (field -> column index, unrecognized header labels). The first
    matching column wins so a duplicated header does not silently shadow the
    earlier one.
    """
    mapping: dict[str, int] = {}
    ignored: list[str] = []
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        if not normalized:
            continue
        for logical, aliases in COLUMN_ALIASES.items():
            if normalized in aliases and logical not in mapping:
                mapping[logical] = index
                break
        else:
            ignored.append(str(header).strip())
    return mapping, ignored


def _cell(row: list, index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decimal(raw: str, field_name: str) -> Decimal:
    # Tolerate values pasted from spreadsheets/POS exports: "$3.50", "3,50",
    # "1,234.00". Rejecting these would fail imports for formatting alone.
    cleaned = raw.replace("$", "").replace("€", "").replace("៛", "").strip()
    if cleaned.count(",") == 1 and "." not in cleaned:
        cleaned = cleaned.replace(",", ".")
    else:
        cleaned = cleaned.replace(",", "")
    try:
        value = Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"{field_name} is not a number: {raw!r}") from exc
    if value < 0:
        raise ValueError(f"{field_name} cannot be negative")
    return value


def _read_csv(payload: bytes) -> list[list[str]]:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = payload.decode("latin-1")
        except UnicodeDecodeError as exc:
            raise ProductImportError("File is not readable text") from exc
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(payload: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ProductImportError(
            "Spreadsheet support is unavailable on the server (openpyxl is not installed)"
        ) from exc
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise ProductImportError(f"Could not read spreadsheet: {exc}") from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_table(filename: str, payload: bytes) -> list[list]:
    """Decode an upload into raw rows based on its extension.

    `.xls` (the pre-2007 binary format) is rejected explicitly rather than
    attempted, since openpyxl only handles the OOXML `.xlsx` family and would
    otherwise fail with a confusing low-level error.
    """
    if len(payload) > MAX_UPLOAD_BYTES:
        raise ProductImportError(
            f"File is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)}MB"
        )
    if not payload.strip():
        raise ProductImportError("File is empty")
    suffix = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    if suffix in {"xlsx", "xlsm"}:
        return _read_xlsx(payload)
    if suffix == "xls":
        raise ProductImportError(
            "Legacy .xls files are not supported. Re-save the file as .xlsx or .csv."
        )
    if suffix in {"csv", "tsv", "txt", ""}:
        return _read_csv(payload)
    raise ProductImportError(f"Unsupported file type: .{suffix}")


def parse_products(filename: str, payload: bytes) -> ImportPreview:
    """Parse an upload into validated rows plus per-row errors.

    Row-level problems are collected rather than raised, so a user can fix a
    few bad rows instead of losing the whole import. Only file-level problems
    (unreadable, missing required columns) raise ProductImportError.
    """
    table = read_table(filename, payload)
    if not table:
        raise ProductImportError("File is empty")

    header_index = next(
        (i for i, row in enumerate(table) if any(_cell(row, c) for c in range(len(row)))),
        None,
    )
    if header_index is None:
        raise ProductImportError("File has no header row")

    headers = [_cell(table[header_index], i) for i in range(len(table[header_index]))]
    mapping, ignored = map_columns(headers)
    missing = [column for column in REQUIRED_COLUMNS if column not in mapping]
    if missing:
        raise ProductImportError(
            "Missing required column(s): "
            + ", ".join(missing)
            + ". Found: "
            + (", ".join(h for h in headers if h) or "nothing")
        )

    preview = ImportPreview(
        columns={logical: headers[index] for logical, index in mapping.items()},
        ignored_columns=ignored,
    )
    seen_skus: set[str] = set()

    for offset, raw in enumerate(table[header_index + 1:], start=header_index + 2):
        if not any(_cell(raw, i) for i in range(len(raw))):
            continue
        if len(preview.rows) + len(preview.errors) >= MAX_ROWS:
            preview.errors.append({
                "row": offset,
                "message": f"Stopped after {MAX_ROWS} rows; split the file and import again.",
            })
            break

        name = _cell(raw, mapping.get("name"))
        if not name:
            preview.errors.append({"row": offset, "message": "Missing product name"})
            continue
        price_raw = _cell(raw, mapping.get("price"))
        if not price_raw:
            preview.errors.append({"row": offset, "message": "Missing price"})
            continue
        try:
            price = _decimal(price_raw, "Price")
        except ValueError as exc:
            preview.errors.append({"row": offset, "message": str(exc)})
            continue

        compare_at_price = None
        compare_raw = _cell(raw, mapping.get("compare_at_price"))
        if compare_raw:
            try:
                compare_at_price = _decimal(compare_raw, "Compare-at price")
            except ValueError as exc:
                preview.errors.append({"row": offset, "message": str(exc)})
                continue
            if compare_at_price < price:
                preview.errors.append({
                    "row": offset,
                    "message": "Compare-at price is lower than price",
                })
                continue

        quantity = None
        quantity_raw = _cell(raw, mapping.get("quantity"))
        if quantity_raw:
            try:
                quantity = int(_decimal(quantity_raw, "Quantity"))
            except ValueError as exc:
                preview.errors.append({"row": offset, "message": str(exc)})
                continue

        sku = _cell(raw, mapping.get("sku")).upper() or None
        if sku and sku in seen_skus:
            preview.errors.append({
                "row": offset,
                "message": f"Duplicate SKU in this file: {sku}",
            })
            continue
        if sku:
            seen_skus.add(sku)

        preview.rows.append(ImportRow(
            row_number=offset,
            name=name,
            price=price,
            sku=sku,
            compare_at_price=compare_at_price,
            category=_cell(raw, mapping.get("category")).lower() or None,
            description=_cell(raw, mapping.get("description")) or None,
            quantity=quantity,
            variant_title=_cell(raw, mapping.get("variant_title")) or None,
        ))

    return preview
